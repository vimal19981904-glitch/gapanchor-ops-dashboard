import os
import sys
import re
import datetime
import json

EXCEL_PATH = r"C:\Users\ARUL XAVIER\OneDrive - gapanchor\Desktop\DemoEnquiry_Extracted.xlsx"
EXCEL_PATH_FALLBACK = r"C:\Users\ARUL XAVIER\OneDrive - gapanchor\Desktop\DemoEnquiry_Extracted_Live.xlsx"

def normalize_country(phone_str, geocoder_country=""):
    cleaned = re.sub(r'[^\d+]', '', phone_str or '')
    if cleaned.count('+') > 1:
        cleaned = '+' + re.sub(r'\+', '', cleaned)
    
    if cleaned.startswith('+91') or (len(cleaned) == 10 and cleaned and cleaned[0] in '6789') or (len(cleaned) == 12 and cleaned.startswith('91')):
        return "India"
    if cleaned.startswith('+1') or (len(cleaned) == 11 and cleaned and cleaned[0] == '1'):
        return "USA"
    if cleaned.startswith('+44'):
        return "United Kingdom"
    if cleaned.startswith('+971'):
        return "United Arab Emirates"
    if cleaned.startswith('+966'):
        return "Saudi Arabia"
    if cleaned.startswith('+33'):
        return "France"
    if cleaned.startswith('+49'):
        return "Germany"
    if cleaned.startswith('+32'):
        return "Belgium"
    if cleaned.startswith('+31'):
        return "Netherlands"
    if cleaned.startswith('+34'):
        return "Spain"
    if cleaned.startswith('+62'):
        return "Indonesia"
    if cleaned.startswith('+52'):
        return "Mexico"
    
    geo = (geocoder_country or "").strip().lower()
    if not geo or geo == "unknown" or geo == "n/a":
        return "India"

    if any(kw in geo for kw in ["india", "karnataka", "gujarat", "madhya pradesh", "maharashtra", "uttar pradesh", "bangalore", "ahmedabad", "delhi", "mumbai", "punjab", "haryana", "kerala", "tamil nadu", "telangana", "andhra", "rajasthan", "baghpat", "baraut", "gwalior", "kalyan", "gundlupet", "malhargarh"]):
        return "India"
    if any(kw in geo for kw in ["united states", "usa", "us", "kansas", "missouri", "california", "north carolina", "ohio", "new hampshire", "michigan", "illinois", "massachusetts", "texas", "new york", "georgia", "pennsylvania", "florida"]):
        return "USA"
    if any(kw in geo for kw in ["canada", "alberta", "ontario", "toronto", "nova scotia", "prince edward"]):
        return "Canada"
    if any(kw in geo for kw in ["mexico", "gomez farias", "sayula", "jal"]):
        return "Mexico"
    if any(kw in geo for kw in ["united kingdom", "uk", "england", "scotland", "wales"]):
        return "United Kingdom"
        
    return geocoder_country.strip() if geocoder_country else "India"

def extract_enquiries():
    try:
        import win32com.client
        import openpyxl
    except ImportError as e:
        print(json.dumps({"success": False, "error": f"Missing required python library: {e}"}))
        return

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mapi = outlook.GetNamespace("MAPI")

        def find_folder(folder, target_name):
            try:
                for subfolder in folder.Folders:
                    if subfolder.Name.lower() == target_name.lower():
                        return subfolder
                    res = find_folder(subfolder, target_name)
                    if res:
                        return res
            except Exception:
                pass
            return None

        demo_folder = None
        try:
            default_root = mapi.GetDefaultFolder(6).Parent
            demo_folder = find_folder(default_root, "Demo enquiry!")
        except Exception:
            pass

        if not demo_folder:
            for folder in mapi.Folders:
                demo_folder = find_folder(folder, "Demo enquiry!")
                if demo_folder:
                    break

        if not demo_folder:
            print(json.dumps({"success": False, "error": "Outlook folder 'Demo enquiry!' not found"}))
            return

        items = demo_folder.Items
        total_items = len(items)
        raw_list = []

        for i in range(1, total_items + 1):
            try:
                email = items.Item(i)
                body = getattr(email, 'Body', '') or ''
                subject = getattr(email, 'Subject', '') or ''
                received_dt = getattr(email, 'ReceivedTime', None)

                received_str = ""
                sort_ts = 0
                if received_dt:
                    try:
                        received_str = received_dt.strftime("%d/%m/%Y %H:%M")
                        sort_ts = received_dt.timestamp()
                    except Exception:
                        received_str = str(received_dt)

                sender_name = getattr(email, 'SenderName', '') or ''
                sender_email = getattr(email, 'SenderEmailAddress', '') or ''

                # Regex matches
                name_match = re.search(r'Name:\s*(.+?)(?:\r?\n|Email:|$)', body, re.IGNORECASE)
                email_match = re.search(r'Email:\s*(.+?)(?:\r?\n|Phone:|$)', body, re.IGNORECASE)
                phone_match = re.search(r'Phone:\s*(.+?)(?:\r?\n|Service\s*Type:|Message:|$)', body, re.IGNORECASE)
                service_match = re.search(r'Service\s*Type:\s*(.+?)(?:\r?\n|Training\s*Type:|Message:|$)', body, re.IGNORECASE)
                training_match = re.search(r'Training\s*Type:\s*(.+?)(?:\r?\n|Message:|$)', body, re.IGNORECASE)
                message_match = re.search(r'Message:\s*(.+?)(?:\s*---|\r?\n|Submitted at:|$)', body, re.IGNORECASE | re.DOTALL)
                country_match = re.search(r'Country:\s*(.+?)(?:\r?\n|$)', body, re.IGNORECASE)
                name = name_match.group(1).strip() if name_match else (sender_name if sender_name else "Unknown")
                email_addr = email_match.group(1).strip() if email_match else (sender_email if "@" in sender_email else "")
                phone = phone_match.group(1).strip() if phone_match else ""
                service_type = service_match.group(1).strip() if service_match else "Training"
                training_type_raw = training_match.group(1).strip() if training_match else ""
                message_raw = message_match.group(1).strip() if message_match else ""
                raw_country = country_match.group(1).strip() if country_match else "India"
                country = normalize_country(phone, raw_country)
                submitted_at = submitted_match.group(1).strip() if submitted_match else received_str

                if "--- Submitted at:" in message_raw:
                    message_raw = message_raw.split("--- Submitted at:")[0].strip()

                training_type = training_type_raw
                combined_text = f"{training_type_raw} {message_raw} {subject}".lower()

                if "proactive" in combined_text:
                    training_type = "Manhattan ProActive"
                elif "manhattan wms" in combined_text or ("manhattan" in combined_text and "wms" in combined_text):
                    training_type = "Manhattan WMS"
                elif "blue yonder" in combined_text or "jda" in combined_text:
                    training_type = "Blue Yonder WMS (JDA)"
                elif "kinaxis" in combined_text:
                    training_type = "Kinaxis RapidResponse"
                elif "sap" in combined_text:
                    training_type = "SAP S/4HANA"
                elif not training_type or training_type.lower() == "training" or training_type.lower() == "n/a":
                    if message_raw and message_raw.lower() != "n/a":
                        training_type = message_raw
                    else:
                        training_type = "General Training"

                message = message_raw if message_raw else training_type

                raw_list.append({
                    "sort_ts": sort_ts,
                    "name": name,
                    "email": email_addr,
                    "phone": phone,
                    "service_type": service_type,
                    "training_type": training_type,
                    "message": message,
                    "country": country,
                    "date_submitted": submitted_at,
                    "received_date": received_str,
                })
            except Exception:
                continue

        # Sort newest first in Python memory
        raw_list.sort(key=lambda x: x["sort_ts"], reverse=True)

        enquiries = raw_list

        # Save to Excel (with lock handling)
        saved_file = EXCEL_PATH
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH) if os.path.exists(EXCEL_PATH) else openpyxl.Workbook()
            if "Sheet 1" in wb.sheetnames:
                ws1 = wb["Sheet 1"]
                ws1.delete_rows(1, ws1.max_row + 1)
            else:
                ws1 = wb.create_sheet("Sheet 1")

            headers = ["Name", "Email ID", "Phone", "Service Type", "Training Type", "Message", "Country", "Date Submitted", "Received Date"]
            ws1.append(headers)

            for item in enquiries:
                ws1.append([
                    item["name"],
                    item["email"],
                    item["phone"],
                    item["service_type"],
                    item["training_type"],
                    item["message"],
                    item["country"],
                    item["date_submitted"],
                    item["received_date"]
                ])

            if "Dashboard" in wb.sheetnames:
                ws_dash = wb["Dashboard"]
                unique_countries_count = len(set(e["country"] for e in enquiries if e["country"]))
                ws_dash.cell(row=6, column=2, value=len(enquiries))
                ws_dash.cell(row=6, column=3, value=unique_countries_count)

            wb.save(EXCEL_PATH)
        except PermissionError:
            saved_file = EXCEL_PATH_FALLBACK
            try:
                wb.save(EXCEL_PATH_FALLBACK)
            except Exception:
                pass

        print(json.dumps({
            "success": True,
            "count": len(enquiries),
            "filePath": saved_file,
            "message": f"Successfully extracted {len(enquiries)} emails from Outlook into {saved_file}"
        }))

    except Exception as err:
        print(json.dumps({"success": False, "error": str(err)}))

if __name__ == "__main__":
    extract_enquiries()
